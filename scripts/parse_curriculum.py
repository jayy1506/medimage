"""
Curriculum and EQB parser module.
Parses MTIP_Curriculum_Data.md and EQB_MTIP-26_updated.xlsx.
"""

import re
import openpyxl
from typing import Dict, List, Any, Tuple


def parse_curriculum_md(md_path: str) -> Dict[str, Any]:
    with open(md_path, "r", encoding="utf-8") as f:
        text = f.read()

    # 1. Sources Register (Section 2)
    sources = []
    sec2_match = re.search(r"## 2\. Source register(.*?)(?=## 3\.)", text, re.DOTALL)
    if sec2_match:
        src_lines = [l.strip() for l in sec2_match.group(1).strip().split("\n") if l.strip().startswith("- **")]
        for l in src_lines:
            m = re.match(r"- \*\*([A-Za-z0-9]+)\*\*:\s*(.*)", l)
            if m:
                code = m.group(1)
                desc = m.group(2)
                sources.append({"code": code, "description": desc})

    # 2. Course Outcomes (Section 3)
    course_outcomes = [
        {"code": "12605.1.1", "short_code": "CO1", "description": "Describe various modalities of medical image sources."},
        {"code": "12605.1.2", "short_code": "CO2", "description": "Explain image processing operations and representation."},
        {"code": "12605.1.3", "short_code": "CO3", "description": "Apply image transformation tools to various applications."},
        {"code": "12605.1.4", "short_code": "CO4", "description": "Design tools to read, visualize and process medical images."}
    ]

    # 3. Decisions (Section 4)
    decisions = []
    sec4_match = re.search(r"## 4\. Data checks and decisions needed(.*?)(?=## 5\.)", text, re.DOTALL)
    if sec4_match:
        dec_matches = re.findall(r"^(\d+)\.\s+(.*?)(?=\n\d+\.|\Z)", sec4_match.group(1).strip(), re.DOTALL | re.MULTILINE)
        for num, desc in dec_matches:
            decisions.append({
                "decision_number": int(num),
                "title": desc.strip().split("\n")[0][:100],
                "description": desc.strip(),
                "status": "open"
            })

    # 4. Source Gaps (Section 5)
    sec5_match = re.search(r"## 5\. Source gaps(.*?)(?=## 6\.)", text, re.DOTALL)
    gaps_text = sec5_match.group(1).strip() if sec5_match else ""

    # 5. Overlap Groups (Section 6)
    overlap_groups = []
    sec6_match = re.search(r"## 6\. Concepts repeated across modules(.*?)(?=## 7\.)", text, re.DOTALL)
    if sec6_match:
        grp_matches = re.findall(r"- \*\*([Gg]\d+)\*\*\s+([^:]+):\s*(.*)", sec6_match.group(1))
        for g_code, g_title, g_members in grp_matches:
            overlap_groups.append({
                "group_code": g_code.upper(),
                "title": g_title.strip(),
                "members_text": g_members.strip()
            })

    # 6. Modules 1 to 6
    modules = {}
    mod_pattern = re.compile(r"^## (\d+)\. Module (\d): ([^\n]+)\n(.*?)(?=(?:^## \d+\.|\Z))", re.DOTALL | re.MULTILINE)
    for m in mod_pattern.finditer(text):
        sec_num = int(m.group(1))
        mod_num = int(m.group(2))
        mod_title = m.group(3).strip()
        mod_body = m.group(4)

        # Header info e.g. "8 hours · EQB CO1 · 40 questions (B15 S15 L10)..."
        hours = 8 if mod_num in (1, 2, 3) else 7
        eqb_co = f"CO{mod_num}" if mod_num <= 5 else "CO5"

        topics = []
        topic_pattern = re.compile(
            r"^### (" + str(mod_num) + r"\.\d+)\s+([^\n]+)\n(.*?)(?=(?:^### " + str(mod_num) + r"\.|\Z))",
            re.DOTALL | re.MULTILINE
        )
        for t_match in topic_pattern.finditer(mod_body):
            t_code = t_match.group(1)
            t_title = t_match.group(2).strip()
            t_body = t_match.group(3).strip()

            # Key terms
            kt_match = re.search(r"-\s*Key terms:\s*(.*?)(?:\n|$)", t_body)
            key_terms = [k.strip() for k in kt_match.group(1).split(",")] if kt_match else []

            # Sources
            src_match = re.search(r"-\s*Sources:\s*(.*?)(?:\n|$)", t_body)
            sources_str = src_match.group(1) if src_match else ""

            # Coverage
            cov_match = re.search(r"Coverage:\s*([^.\n]+)", t_body)
            cov_str = cov_match.group(1).strip() if cov_match else "none"
            if "none" in cov_str.lower():
                coverage = "none"
                status = "needs_source"
            elif "partial" in cov_str.lower():
                coverage = "partial"
                status = "skeleton"
            elif "full" in cov_str.lower():
                coverage = "full"
                status = "skeleton"
            else:
                coverage = "none"
                status = "needs_source"

            # EQB question tokens
            eqb_match = re.search(r"-\s*EQB:\s*(.*?)(?:\n|$)", t_body)
            eqb_str = eqb_match.group(1) if eqb_match else ""
            count_match = re.search(r"\((\d+)\)", eqb_str)
            doc_count = int(count_match.group(1)) if count_match else 0
            clean_str = re.sub(r"\(\d+\)", "", eqb_str)
            q_tokens = re.findall(r"\b([BSL]\d{2})\b", clean_str)

            topics.append({
                "code": t_code,
                "title": t_title,
                "key_terms": key_terms,
                "sources_str": sources_str,
                "coverage": coverage,
                "status": status,
                "doc_count": doc_count,
                "q_tokens": q_tokens
            })

        modules[mod_num] = {
            "module_number": mod_num,
            "title": mod_title,
            "hours": hours,
            "eqb_co": eqb_co,
            "topics": topics
        }

    return {
        "sources": sources,
        "course_outcomes": course_outcomes,
        "decisions": decisions,
        "source_gaps_text": gaps_text,
        "overlap_groups": overlap_groups,
        "modules": modules
    }


def parse_eqb_excel(xlsx_path: str) -> Dict[Tuple[int, str], List[Dict[str, Any]]]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["MTIP"]

    # Header is at row 3, data from row 4
    # Columns:
    # 0: S. No.
    # 1: Question
    # 2: Ref Book Name
    # 3: Page No
    # 4: Distribution Level (MK / DK)
    # 5: Difficulty Level/ Bloom's Level
    # 6: Marks
    # 7: Subject
    # 8: Paper Type 1 or 2
    # 9: Topic (e.g. 24.1-Module1)
    # 10: Sub Topic
    # 11: Question Type (BAQ / SAQ / LAQ)
    # 12: Reference Details
    # 13: ModelAnswer
    # 14: Model ImagePath
    # 15: Competency / CO

    records_by_module_type = {}
    for r_idx in range(4, ws.max_row + 1):
        vals = [cell.value for cell in ws[r_idx]]
        if not any(vals):
            continue
        topic_val = str(vals[9] or "").strip()
        m_match = re.search(r"Module\s*(\d+)", topic_val, re.IGNORECASE)
        if not m_match:
            continue
        mod_num = int(m_match.group(1))
        qtype = str(vals[11] or "").strip()

        key = (mod_num, qtype)
        if key not in records_by_module_type:
            records_by_module_type[key] = []

        records_by_module_type[key].append({
            "eqb_row": r_idx,
            "s_no": vals[0],
            "original_text": vals[1],
            "ref_book": vals[2],
            "page_no": vals[3],
            "distribution_level": str(vals[4] or "").strip(),
            "bloom": int(vals[5]) if vals[5] is not None and str(vals[5]).isdigit() else 1,
            "marks": int(vals[6]) if vals[6] is not None and str(vals[6]).isdigit() else (2 if qtype == "BAQ" else (4 if qtype == "SAQ" else 10)),
            "subject": vals[7],
            "topic_str": vals[9],
            "sub_topic": vals[10],
            "qtype": qtype,
            "eqb_co": int(vals[15]) if vals[15] is not None and str(vals[15]).isdigit() else mod_num
        })

    return records_by_module_type


def get_module_questions(mod_num: int, cur_data: Dict[str, Any], eqb_data: Dict[Tuple[int, str], List[Dict[str, Any]]]) -> List[Dict[str, Any]]:
    type_code_map = {"BAQ": "B", "SAQ": "S", "LAQ": "L"}
    inv_code_map = {"B": "BAQ", "S": "SAQ", "L": "LAQ"}

    # Map q_token -> topic_code
    mod_info = cur_data["modules"][mod_num]
    q_to_topic = {}
    for t in mod_info["topics"]:
        for tok in t["q_tokens"]:
            full_id = f"MTIP-M{mod_num}-{tok}"
            q_to_topic[full_id] = t["code"]

    all_questions = []
    for qtype in ["BAQ", "SAQ", "LAQ"]:
        prefix = type_code_map[qtype]
        rows = eqb_data.get((mod_num, qtype), [])
        for idx, row in enumerate(rows, 1):
            q_id = f"MTIP-M{mod_num}-{prefix}{idx:02d}"
            assigned_topic = q_to_topic.get(q_id)
            all_questions.append({
                "id": q_id,
                "module": mod_num,
                "topic_code": assigned_topic,
                "type": qtype,
                "marks": row["marks"],
                "bloom": row["bloom"],
                "distribution_level": row["distribution_level"],
                "eqb_co": row["eqb_co"],
                "eqb_row": row["eqb_row"],
                "original_text": row["original_text"],
                "status": "draft"
            })

    return all_questions
